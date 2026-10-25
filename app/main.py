"""Local-first web API. Uploaded source documents are processed in memory, not persisted."""
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from xml.sax.saxutils import escape

from .agent import Handover, generate_handover

BASE = Path(__file__).resolve().parent.parent
app = FastAPI(title="ExitKit AI", version="0.1.0")
app.mount("/static", StaticFiles(directory=BASE / "static"), name="static")
MAX_BYTES = 5 * 1024 * 1024


@app.get("/")
async def index():
    return FileResponse(BASE / "static" / "index.html")


@app.get("/api/health")
async def health():
    return {"status": "ok"}


def read_document(name: str, content: bytes) -> str:
    extension = Path(name).suffix.lower()
    try:
        if extension in (".txt", ".md"):
            return content.decode("utf-8-sig")
        if extension == ".pdf":
            return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
        if extension == ".docx":
            doc = Document(BytesIO(content))
            blocks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    blocks.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(blocks)
    except Exception as err:
        raise HTTPException(422, "Не удалось прочитать файл. Проверьте формат и повреждения.") from err
    raise HTTPException(415, "Поддерживаются только .txt, .md, .pdf, .docx")


@app.post("/api/generate", response_model=Handover)
async def generate(notes: str = Form(""), files: list[UploadFile] = File(default=[])):
    chunks = [notes.strip()]
    for item in files:
        raw = await item.read(MAX_BYTES + 1)
        if len(raw) > MAX_BYTES:
            raise HTTPException(413, "Каждый файл должен быть не более 5 МБ")
        chunks.append(read_document(item.filename or "", raw))
    merged = "\n\n".join(c for c in chunks if c.strip())
    if not merged:
        raise HTTPException(400, "Добавьте текст или файл")
    if len(merged) > 80000:
        raise HTTPException(413, "Слишком много текста: лимит 80 000 символов")
    try:
        return await generate_handover(merged)
    except Exception as err:
        raise HTTPException(502, "AI-модель недоступна или ответила некорректно. Проверьте LLM_BASE_URL и LLM_MODEL.") from err


def get_font() -> str:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for item in candidates:
        if Path(item).exists():
            pdfmetrics.registerFont(TTFont("ExitKitUnicode", item))
            return "ExitKitUnicode"
    raise RuntimeError("Для кириллицы установите шрифт DejaVu Sans (пакет fonts-dejavu-core)")


def pdf_content(payload: Handover) -> bytes:
    font = get_font()
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=(595, 842), rightMargin=44, leftMargin=44, topMargin=44, bottomMargin=44)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="EKTitle", fontName=font, fontSize=19, leading=25, spaceAfter=18, textColor=colors.HexColor("#14243b")))
    styles.add(ParagraphStyle(name="EKHead", fontName=font, fontSize=12, leading=17, spaceBefore=13, spaceAfter=8, textColor=colors.HexColor("#185c80")))
    styles.add(ParagraphStyle(name="EKBody", fontName=font, fontSize=9, leading=14, spaceAfter=6))
    story = [Paragraph("ExitKit AI — пакет передачи дел", styles["EKTitle"])]
    def para(text):
        story.append(Paragraph(escape(str(text)).replace("\n", "<br/>"), styles["EKBody"]))
    def section(label):
        story.append(Paragraph(label, styles["EKHead"]))
    if payload.role:
        para("Роль: " + payload.role)
    section("Обзор")
    para(payload.summary or "Нет описания")
    section("Текущие задачи")
    if not payload.tasks:
        para("Задачи не выделены")
    for n, task in enumerate(payload.tasks, 1):
        para(f"{n}. {task.title} | статус: {task.status} | приоритет: {task.priority} | ответственный: {task.owner} | срок: {task.deadline or 'не указан'}")
        for step in task.next_steps:
            para("Следующий шаг: " + step)
        for blocker in task.blockers:
            para("Блокер: " + blocker)
        if task.source_excerpt:
            para("Источник: «" + task.source_excerpt + "»")
    for title, items in [("Открытые вопросы", payload.open_questions), ("Контакты", payload.contacts), ("Чек-лист передачи", payload.checklist), ("Что нужно проверить", payload.warnings)]:
        section(title)
        if not items:
            para("Нет данных")
        for index, item in enumerate(items, 1):
            para(f"{index}. {item}")
    para("Сгенерировано автоматически. Требуется проверка сотрудником и менеджером.")
    doc.build(story)
    return stream.getvalue()


@app.post("/api/export/pdf")
async def export_pdf(payload: Handover):
    try:
        content = pdf_content(payload)
    except RuntimeError as err:
        raise HTTPException(500, str(err)) from err
    return Response(content, media_type="application/pdf", headers={"Content-Disposition": 'attachment; filename="exitkit-handover.pdf"'})
